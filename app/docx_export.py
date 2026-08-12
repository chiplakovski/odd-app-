"""Word (.docx) inspection report generation from grouped work items."""
from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from .config import MASTER_TEMPLATE, ProjectInfo
from .grouping import group_items
from .models import WorkItem

_METHOD_LABELS = {
    "Visual inspection": "Visual inspection",
    "Visual & NDT control": "Visual inspection, NDT control",
}


def _clear_paragraph(paragraph) -> None:
    p = paragraph._element
    for child in list(p):
        p.remove(child)


def set_cell_text(cell, text: str, size: float = 8.0, bold: bool = False,
                   alignment=None, font_name: str = "Arial") -> None:
    # Keep cell/table geometry but replace visible content cleanly.
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    _clear_paragraph(p)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if alignment is not None:
        p.alignment = alignment
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size)
    run.bold = bold


def _compact_summary(text: str, max_chars: int) -> str:
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    if not lines:
        return "Work completed as specified in the work list."
    out: list[str] = []
    used = 0
    for line in lines:
        cleaned = re.sub(r"^[-–—_]{4,}\s*", "", line).strip()
        if not cleaned:
            continue
        addition = len(cleaned) + (1 if out else 0)
        if out and used + addition > max_chars:
            break
        if not out and len(cleaned) > max_chars:
            cleaned = cleaned[: max_chars - 3].rsplit(" ", 1)[0] + "..."
        out.append(cleaned)
        used += addition
    compact = "; ".join(out)
    if len(compact) < len(text.strip()) and not compact.endswith("..."):
        compact += "; ..."
    return compact


def compose_group_description(group: list[WorkItem]) -> str:
    # Allocate space to every litra so no job number appears without a description.
    total_budget = 1750
    label_allowance = sum(len(f"Item {item.number}: ") + 2 for item in group)
    per_item = max(150, min(620, (total_budget - label_allowance) // max(1, len(group))))
    blocks: list[str] = []
    for item in group:
        summary = item.summary.strip() or item.description.strip()
        blocks.append(f"Item {item.number}: {_compact_summary(summary, per_item)}")
    return "\n".join(blocks)


def _results_and_supplement(info: ProjectInfo, group: list[WorkItem]) -> tuple[str, str]:
    result_mode = (info.completion_result or "Automatic (from work list)").strip()
    if result_mode == "Final inspection":
        return (
            "Final inspection completed. No remarks, Accepted.",
            "Final inspection has been carried out and the work delivered in good order.",
        )
    if result_mode == "Completed jobs":
        return (
            "Completed jobs. No remarks, Accepted.",
            "Completed jobs have been inspected and delivered in good order.",
        )
    if result_mode == "Partially completed":
        return (
            "Partially completed. Final inspection pending completion.",
            "Completed work has been inspected. Final approval is pending completion of the remaining work.",
        )
    all_finished = all(i.status.lower() in {"finished", "completed"} for i in group)
    if all_finished:
        return "No remarks, Accepted", "Inspections have been done and delivered in good order."
    return (
        "Work in progress, final inspection pending completion.",
        "Completed work has been visually inspected. Final approval is pending completion.",
    )


def fill_report_table(table, group: list[WorkItem], info: ProjectInfo) -> None:
    content = table.cell(2, 0)
    if len(content.tables) < 5:
        raise RuntimeError("The master template structure is not valid.")
    project = content.tables[0]
    set_cell_text(project.cell(0, 1), info.project_name, 8.5)
    set_cell_text(project.cell(1, 1), info.project_number, 8.5)
    jobs = "/".join(str(i.number) for i in group)
    job_size = 8.0 if len(jobs) <= 22 else 7.0 if len(jobs) <= 35 else 6.2
    set_cell_text(project.cell(2, 1), jobs, job_size)

    desc = compose_group_description(group)
    desc_size = 8.0
    if len(group) >= 4 or len(desc) > 750:
        desc_size = 7.0
    if len(group) >= 5 or len(desc) > 1250:
        desc_size = 6.4
    set_cell_text(content.tables[1].cell(0, 0), desc, desc_size)

    method_text = _METHOD_LABELS.get(info.inspection_method, info.inspection_method or "Visual inspection")
    set_cell_text(content.tables[2].cell(0, 0), method_text, 8.0)

    results, supplement = _results_and_supplement(info, group)
    set_cell_text(content.tables[3].cell(0, 0), results, 8.0)
    set_cell_text(content.tables[4].cell(0, 0), supplement, 8.0)

    set_cell_text(table.cell(3, 2), info.inspection_company, 7.0)
    set_cell_text(table.cell(3, 4), info.inspector_names, 6.6)
    set_cell_text(table.cell(3, 6), info.report_date, 7.0)
    set_cell_text(table.cell(4, 2), info.supervisor_company, 7.0)
    set_cell_text(table.cell(4, 4), info.supervisor_name, 6.6)


def _page_break_paragraph() -> OxmlElement:
    # A one-twip exact-height paragraph, matching the approved source document. A normal Word
    # paragraph is too tall and can create a blank page after a full-height form.
    page_p = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "1")
    spacing.set(qn("w:lineRule"), "exact")
    p_pr.append(spacing)
    page_p.append(p_pr)
    run = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run.append(br)
    page_p.append(run)
    return page_p


def generate_docx(items: list[WorkItem], info: ProjectInfo, output_path: Path,
                   master_template: Path = MASTER_TEMPLATE) -> int:
    groups = group_items(items, only_finished=info.only_finished)
    if not groups:
        raise ValueError("No included items match the selected status filter.")
    if not master_template.exists():
        raise FileNotFoundError(f"Master template not found: {master_template}")

    doc = Document(str(master_template))
    base_table = doc.tables[0]
    fill_report_table(base_table, groups[0], info)

    for group in groups[1:]:
        doc._body._body.insert(len(doc._body._body) - 1, _page_break_paragraph())
        new_tbl_xml = deepcopy(base_table._tbl)
        doc._body._body.insert(len(doc._body._body) - 1, new_tbl_xml)
        new_table = doc.tables[-1]
        fill_report_table(new_table, group, info)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return len(groups)
